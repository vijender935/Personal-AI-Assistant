"""Document and OCR extraction helpers for the authenticated file pipeline."""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".pdf", ".docx",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
}


def is_supported_document(path: str | Path) -> bool:
    return Path(path).suffix.lower() in SUPPORTED_DOCUMENT_EXTENSIONS


def _text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _csv_file(path: Path) -> str:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    rows = csv.reader(io.StringIO(raw))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _json_file(path: Path) -> str:
    raw = _text_file(path)
    try:
        return json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return raw


def _pdf_file(path: Path) -> str:
    import fitz

    native_pages = []
    needs_ocr = False
    with fitz.open(path) as document:
        for page in document:
            text = page.get_text("text").strip()
            native_pages.append(text)
            if not text:
                needs_ocr = True

    if not needs_ocr:
        return "\n\n".join(page for page in native_pages if page)

    from ocr import ocr_pdf
    return ocr_pdf(path)


def _docx_file(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _image_file(path: Path) -> str:
    from ocr import ocr_image
    return ocr_image(path)


def extract_text(path: str | Path) -> str:
    candidate = Path(path)
    suffix = candidate.suffix.lower()
    if suffix in {".txt", ".md", ".xml"}:
        return _text_file(candidate)
    if suffix == ".csv":
        return _csv_file(candidate)
    if suffix == ".json":
        return _json_file(candidate)
    if suffix == ".pdf":
        return _pdf_file(candidate)
    if suffix == ".docx":
        return _docx_file(candidate)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        return _image_file(candidate)
    raise ValueError(f"Unsupported document type: {suffix or 'unknown'}")


def extract_and_limit(path: str | Path, max_chars: int = 500_000) -> str:
    text = extract_text(path).strip()
    if not text:
        raise ValueError("No readable text was found in the document.")
    return text[:max_chars]
