from document_parser import extract_text, is_supported_document
import document_parser


def test_txt_extraction(tmp_path):
    path=tmp_path/"notes.txt"
    path.write_text("hello\nworld",encoding="utf-8")
    assert "hello" in extract_text(path)


def test_csv_extraction(tmp_path):
    path=tmp_path/"data.csv"
    path.write_text("name,score\nVijender,10\n",encoding="utf-8")
    text=extract_text(path)
    assert "Vijender" in text
    assert "10" in text


def test_json_extraction(tmp_path):
    path=tmp_path/"data.json"
    path.write_text('{"name":"Vijender","score":10}',encoding="utf-8")
    text=extract_text(path)
    assert '"name": "Vijender"' in text
    assert '"score": 10' in text


def test_docx_extraction(tmp_path):
    from docx import Document
    path=tmp_path/"report.docx"
    doc=Document()
    doc.add_paragraph("FastAPI document")
    table=doc.add_table(rows=1,cols=2)
    table.rows[0].cells[0].text="Python"
    table.rows[0].cells[1].text="RAG"
    doc.save(path)
    text=extract_text(path)
    assert "FastAPI document" in text
    assert "Python | RAG" in text


def test_image_ocr_extraction(tmp_path, monkeypatch):
    path=tmp_path/"scan.png"
    path.write_bytes(b"not-a-real-image")
    monkeypatch.setattr(document_parser, "_image_file", lambda _: "OCR text from image")
    assert extract_text(path) == "OCR text from image"


def test_supported_documents():
    assert is_supported_document("report.pdf")
    assert is_supported_document("report.docx")
    assert is_supported_document("data.csv")
    assert is_supported_document("data.json")
    assert is_supported_document("scan.png")
    assert not is_supported_document("video.mp4")
